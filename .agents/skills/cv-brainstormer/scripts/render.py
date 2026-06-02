#!/usr/bin/env python3
"""
CV Brainstormer — Report Renderer
Converts Markdown report to DOCX and PDF.

Usage:
    python render.py <input.md> [--output-dir output/]

Examples:
    python render.py output/final-report.md
    python render.py output/final-report.md --output-dir output/

Dependencies:
    pip install python-docx markdown weasyprint --break-system-packages
"""

import sys
import os
import re
import argparse
from pathlib import Path
from datetime import datetime


def check_dependencies():
    missing = []
    try:
        import docx
    except ImportError:
        missing.append("python-docx")
    try:
        import markdown
    except ImportError:
        missing.append("markdown")
    try:
        import weasyprint
    except ImportError:
        missing.append("weasyprint")

    if missing:
        print("ERROR: Missing dependencies:")
        for dep in missing:
            print(f"  - {dep}")
        print()
        print("Install with:")
        print(f"  pip install {' '.join(missing)} --break-system-packages")
        sys.exit(1)


def extract_cv_section(content: str) -> str:
    """Extract the revised CV section from the full report."""
    # Look for the revised CV section
    patterns = [
        r"## Revised CV Draft.*?\n(.*?)(?=\n---|\n## |\Z)",
        r"## Draft CV yang Telah Diperbaiki.*?\n(.*?)(?=\n---|\n## |\Z)",
    ]

    for pattern in patterns:
        match = re.search(pattern, content, re.DOTALL | re.IGNORECASE)
        if match:
            cv_content = match.group(1).strip()
            # Remove the note/blockquote at the beginning
            cv_content = re.sub(r'^>.*?\n\n', '', cv_content, flags=re.DOTALL)
            return cv_content

    return ""


def md_to_docx(md_content: str, output_path: Path, doc_type: str = "report"):
    """Convert Markdown content to DOCX."""
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # Style defaults
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    lines = md_content.split('\n')
    i = 0

    while i < len(lines):
        line = lines[i]

        # H1
        if line.startswith('# ') and not line.startswith('## '):
            para = doc.add_heading(line[2:].strip(), level=1)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if doc_type == "cv" else WD_ALIGN_PARAGRAPH.LEFT

        # H2
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)

        # H3
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)

        # H4
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=4)

        # Horizontal rule
        elif line.strip() in ('---', '***', '___'):
            doc.add_paragraph('─' * 60)

        # Bullet points
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:]
            # Handle bold in bullets
            para = doc.add_paragraph(style='List Bullet')
            _add_formatted_run(para, text)

        # Numbered list
        elif re.match(r'^\d+\. ', line.strip()):
            text = re.sub(r'^\d+\. ', '', line.strip())
            para = doc.add_paragraph(style='List Number')
            _add_formatted_run(para, text)

        # Blockquote
        elif line.strip().startswith('>'):
            text = line.strip().lstrip('> ')
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.5)
            run = para.add_run(text)
            run.italic = True

        # Table row
        elif line.strip().startswith('|') and '|' in line:
            # Skip separator rows
            if re.match(r'^[\|\-\s:]+$', line.strip()):
                i += 1
                continue
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            if any(cells):
                # Simple table handling — just write as tab-separated
                para = doc.add_paragraph()
                para.add_run('\t'.join(cells))

        # Empty line
        elif line.strip() == '':
            if i > 0 and lines[i-1].strip() != '':
                doc.add_paragraph()

        # Regular paragraph
        else:
            if line.strip():
                para = doc.add_paragraph()
                _add_formatted_run(para, line.strip())

        i += 1

    doc.save(str(output_path))
    print(f"DOCX saved: {output_path}")


def _add_formatted_run(para, text: str):
    """Add a run with bold/italic formatting parsed from markdown."""
    # Simple bold/italic parser
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = para.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = para.add_run(part[1:-1])
            run.font.name = 'Courier New'
        else:
            para.add_run(part)


def md_to_pdf(md_content: str, output_path: Path):
    """Convert Markdown content to PDF via HTML."""
    import markdown
    import weasyprint

    # Convert MD to HTML
    md = markdown.Markdown(extensions=['tables', 'fenced_code', 'nl2br'])
    html_body = md.convert(md_content)

    # Wrap in styled HTML
    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{
    font-family: 'Segoe UI', Calibri, Arial, sans-serif;
    font-size: 11pt;
    line-height: 1.6;
    color: #222;
    max-width: 800px;
    margin: 0 auto;
    padding: 40px;
  }}
  h1 {{ font-size: 20pt; color: #1a1a2e; border-bottom: 2px solid #1a1a2e; padding-bottom: 8px; }}
  h2 {{ font-size: 15pt; color: #16213e; border-bottom: 1px solid #ccc; padding-bottom: 4px; margin-top: 24px; }}
  h3 {{ font-size: 12pt; color: #0f3460; margin-top: 16px; }}
  h4 {{ font-size: 11pt; color: #333; }}
  blockquote {{ border-left: 4px solid #1a1a2e; padding-left: 16px; color: #555; font-style: italic; }}
  code {{ background: #f4f4f4; padding: 2px 6px; border-radius: 3px; font-family: 'Courier New', monospace; font-size: 10pt; }}
  pre {{ background: #f4f4f4; padding: 12px; border-radius: 6px; overflow-x: auto; }}
  table {{ border-collapse: collapse; width: 100%; margin: 12px 0; }}
  th {{ background: #1a1a2e; color: white; padding: 8px 12px; text-align: left; }}
  td {{ border: 1px solid #ddd; padding: 6px 12px; }}
  tr:nth-child(even) {{ background: #f9f9f9; }}
  ul, ol {{ padding-left: 24px; }}
  li {{ margin: 4px 0; }}
  hr {{ border: none; border-top: 1px solid #ddd; margin: 24px 0; }}
  .page-break {{ page-break-before: always; }}
</style>
</head>
<body>
{html_body}
</body>
</html>"""

    weasyprint.HTML(string=html).write_pdf(str(output_path))
    print(f"PDF saved: {output_path}")


def main():
    parser = argparse.ArgumentParser(description='Render CV Brainstormer Markdown to DOCX and PDF')
    parser.add_argument('input', help='Input .md file')
    parser.add_argument('--output-dir', default='output', help='Output directory (default: output/)')
    parser.add_argument('--no-cv', action='store_true', help='Skip CV extraction and rendering')
    args = parser.parse_args()

    check_dependencies()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"ERROR: File not found: {input_path}")
        sys.exit(1)

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    # Read markdown content
    with open(input_path, 'r', encoding='utf-8') as f:
        content = f.read()

    timestamp = datetime.now().strftime('%Y-%m-%d')
    stem = input_path.stem

    # Render full report
    print("\n=== Rendering Full Report ===")
    report_docx = output_dir / f"{stem}.docx"
    report_pdf = output_dir / f"{stem}.pdf"
    md_to_docx(content, report_docx, doc_type="report")
    md_to_pdf(content, report_pdf)

    # Extract and render CV section separately
    if not args.no_cv:
        cv_content = extract_cv_section(content)
        if cv_content:
            print("\n=== Rendering Revised CV ===")
            cv_docx = output_dir / f"cv-revised-{timestamp}.docx"
            cv_pdf = output_dir / f"cv-revised-{timestamp}.pdf"
            md_to_docx(cv_content, cv_docx, doc_type="cv")
            md_to_pdf(cv_content, cv_pdf)

            # Save CV as markdown too
            cv_md = output_dir / f"cv-revised-{timestamp}.md"
            with open(cv_md, 'w', encoding='utf-8') as f:
                f.write(cv_content)
            print(f"CV Markdown saved: {cv_md}")
        else:
            print("Note: No 'Revised CV Draft' section found in report. Skipping CV extraction.")

    print("\n✅ Done! All files saved to:", output_dir)


if __name__ == "__main__":
    main()
