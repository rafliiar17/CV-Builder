"""
Tests for scripts/render_outputs.py.
Verifies converting Markdown with headings, tables, bold/italic runs, and lists
to DOCX and PDF documents.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

# Add repo root to sys.path
REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from scripts.render_outputs import (
    is_table_separator,
    main,
    parse_table_row,
    render_docx,
    render_pdf,
    strip_md,
)

SAMPLE_MARKDOWN = """# John Doe - Curriculum Vitae

Contact: john.doe@example.com | **Jakarta, Indonesia**

## Executive Summary

Experienced **Data Analyst** and *Application Support Specialist* with a proven track record.
Proficient in `SQL`, Python, and database troubleshooting.

## Technical Skills

| Domain | Technologies | Experience Level |
| --- | --- | --- |
| Databases | PostgreSQL, MySQL | **Advanced** |
| Languages | Python, Bash | *Intermediate* |
| Tools | Git, Docker, Linux | Working Knowledge |

## Professional Experience

### Lead Support Analyst - Tech Corp
- Resolved over **150+** critical incident tickets with high satisfaction.
- Automated reporting pipelines using *Python* and scheduled cron jobs.
- Maintained [Documentation](https://example.com) for L1/L2 escalation runbooks.
"""


def test_strip_md_helper():
    assert strip_md("**bold** and *italic*") == "bold and italic"
    assert strip_md("`code` and [link](url)") == "code and link"
    assert strip_md("***bold italic***") == "bold italic"


def test_parse_markdown_table_row():
    row = "| Domain | Technologies | Level |"
    assert parse_table_row(row) == ["Domain", "Technologies", "Level"]

    sep_row = "| --- | :---: | ---: |"
    assert is_table_separator(sep_row) is True

    data_row = "| Databases | PostgreSQL | Advanced |"
    assert is_table_separator(data_row) is False


def test_render_docx_structure(tmp_path: Path):
    md_file = tmp_path / "sample_cv.md"
    md_file.write_text(SAMPLE_MARKDOWN, encoding="utf-8")

    docx_path = render_docx(md_file)
    assert docx_path.exists()
    assert docx_path.suffix == ".docx"
    assert docx_path.stat().st_size > 0

    doc = Document(docx_path)

    # 1. Verify headings exist
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert len(headings) >= 3
    assert "John Doe - Curriculum Vitae" in headings[0]
    assert "Executive Summary" in headings[1]
    assert "Technical Skills" in headings[2]

    # 2. Verify tables exist
    assert len(doc.tables) >= 1
    table = doc.tables[0]

    # Verify table rows and cells
    assert len(table.rows) == 4  # 1 header + 3 data rows
    assert len(table.columns) == 3

    # Verify header cells
    header_cells = [c.text for c in table.rows[0].cells]
    assert header_cells == ["Domain", "Technologies", "Experience Level"]

    # Verify header cell runs have bold formatting
    header_runs = table.rows[0].cells[0].paragraphs[0].runs
    assert len(header_runs) > 0
    assert header_runs[0].bold is True

    # Verify data cell contents
    row1_cells = [c.text for c in table.rows[1].cells]
    assert row1_cells == ["Databases", "PostgreSQL, MySQL", "Advanced"]

    # Verify formatted runs within table cells
    # Level cell in row 1 was "**Advanced**" -> bold run
    adv_runs = table.rows[1].cells[2].paragraphs[0].runs
    assert len(adv_runs) > 0
    assert adv_runs[0].text == "Advanced"
    assert adv_runs[0].bold is True

    # Level cell in row 2 was "*Intermediate*" -> italic run
    int_runs = table.rows[2].cells[2].paragraphs[0].runs
    assert len(int_runs) > 0
    assert int_runs[0].text == "Intermediate"
    assert int_runs[0].italic is True

    # 3. Verify paragraphs and formatted runs outside tables
    all_runs = [run for p in doc.paragraphs for run in p.runs]
    bold_runs = [r.text for r in all_runs if r.bold]
    italic_runs = [r.text for r in all_runs if r.italic]

    assert "Data Analyst" in bold_runs
    assert "Application Support Specialist" in italic_runs


def test_render_pdf_structure(tmp_path: Path):
    md_file = tmp_path / "sample_cv.md"
    md_file.write_text(SAMPLE_MARKDOWN, encoding="utf-8")

    pdf_path = render_pdf(md_file)
    assert pdf_path.exists()
    assert pdf_path.suffix == ".pdf"
    assert pdf_path.stat().st_size > 0

    # Verify PDF magic bytes
    with open(pdf_path, "rb") as f:
        header = f.read(5)
        assert header == b"%PDF-"


def test_render_cli_main(tmp_path: Path):
    md_file = tmp_path / "test_cli.md"
    md_file.write_text(SAMPLE_MARKDOWN, encoding="utf-8")

    # Success invocation
    ret = main([str(md_file)])
    assert ret == 0
    assert (tmp_path / "test_cli.docx").exists()
    assert (tmp_path / "test_cli.pdf").exists()

    # Missing arguments returns usage code 2
    assert main([]) == 2

    # Nonexistent file returns error code 1
    assert main([str(tmp_path / "missing.md")]) == 1


def test_batch_rendering(tmp_path: Path):
    run_dir = tmp_path / "candidate-run"
    (run_dir / "reports").mkdir(parents=True)
    (run_dir / "cv" / "role-a").mkdir(parents=True)
    (run_dir / "input").mkdir(parents=True)
    (run_dir / "scratch").mkdir(parents=True)

    (run_dir / "reports" / "final-report-bilingual.md").write_text(SAMPLE_MARKDOWN, encoding="utf-8")
    (run_dir / "cv" / "role-a" / "cv-test-en.md").write_text(SAMPLE_MARKDOWN, encoding="utf-8")
    (run_dir / "input" / "raw-cv.md").write_text("# Raw Input", encoding="utf-8")
    (run_dir / "scratch" / "notes.md").write_text("# Scratch Notes", encoding="utf-8")

    # Test --run-dir
    ret = main(["--run-dir", str(run_dir)])
    assert ret == 0

    assert (run_dir / "reports" / "final-report-bilingual.docx").exists()
    assert (run_dir / "reports" / "final-report-bilingual.pdf").exists()
    assert (run_dir / "cv" / "role-a" / "cv-test-en.docx").exists()
    assert (run_dir / "cv" / "role-a" / "cv-test-en.pdf").exists()

    # Verify input and scratch files are skipped
    assert not (run_dir / "input" / "raw-cv.docx").exists()
    assert not (run_dir / "scratch" / "notes.docx").exists()
