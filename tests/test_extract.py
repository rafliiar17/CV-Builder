"""
Tests for scripts/extract.py.
Verifies plain text extraction, DOCX extraction, and error handling
for nonexistent and unsupported file formats.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from docx import Document

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from scripts.extract import (
    extract_docx,
    extract_file,
    extract_txt,
    main,
)


def test_extract_txt_plain_file(tmp_path: Path):
    """Test extracting from a plain .txt file."""
    txt_file = tmp_path / "sample_cv.txt"
    sample_text = "Jane Doe\nSoftware Engineer\n5 years experience in Python."
    txt_file.write_text(sample_text, encoding="utf-8")

    result = extract_txt(txt_file)
    assert result == sample_text


def test_extract_txt_markdown_file(tmp_path: Path):
    """Test extracting from a markdown .md file."""
    md_file = tmp_path / "cv.md"
    sample_md = "# Candidate\n- Skill 1\n- Skill 2"
    md_file.write_text(sample_md, encoding="utf-8")

    result = extract_txt(md_file)
    assert result == sample_md


def test_extract_docx_paragraphs_and_tables(tmp_path: Path):
    """Test extracting text and tables from a docx file."""
    docx_file = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("First Paragraph: Work Experience")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"
    table.rows[1].cells[0].text = "Val A"
    table.rows[1].cells[1].text = "Val B"
    doc.save(docx_file)

    extracted = extract_docx(docx_file)
    assert "First Paragraph: Work Experience" in extracted
    assert "Header A | Header B" in extracted
    assert "Val A | Val B" in extracted


def test_extract_file_wrapper(tmp_path: Path):
    """Test extract_file writes structured output file."""
    txt_file = tmp_path / "input.txt"
    txt_file.write_text("Hello Extraction", encoding="utf-8")

    out_file = tmp_path / "output" / "extracted.txt"
    result = extract_file(txt_file, out_file)

    assert result == "Hello Extraction"
    assert out_file.exists()
    content = out_file.read_text(encoding="utf-8")
    assert "=== EXTRACTED FROM: input.txt ===" in content
    assert "Hello Extraction" in content
    assert "=== END OF EXTRACTION ===" in content


def test_extract_nonexistent_file_raises_error(tmp_path: Path):
    """Test nonexistent file raises FileNotFoundError."""
    missing_file = tmp_path / "does_not_exist.txt"
    with pytest.raises(FileNotFoundError):
        extract_txt(missing_file)

    with pytest.raises(FileNotFoundError):
        extract_file(missing_file)


def test_extract_invalid_format_raises_error(tmp_path: Path):
    """Test unsupported format raises ValueError."""
    bad_file = tmp_path / "picture.png"
    bad_file.write_text("fake binary", encoding="utf-8")

    with pytest.raises(ValueError) as excinfo:
        extract_file(bad_file)
    assert "Unsupported file" in str(excinfo.value)
    assert ".png" in str(excinfo.value)


def test_extract_cli_main(tmp_path: Path, capsys):
    """Test CLI main() behavior and exit codes."""
    # 1. No arguments -> error code 1
    assert main([]) == 1

    # 2. Nonexistent file -> error code 1
    assert main([str(tmp_path / "missing.txt")]) == 1

    # 3. Unsupported format -> error code 1
    bad_format = tmp_path / "file.xyz"
    bad_format.write_text("data", encoding="utf-8")
    assert main([str(bad_format)]) == 1

    # 4. Successful extraction -> exit code 0
    valid_file = tmp_path / "good.txt"
    valid_file.write_text("Sample Good CV", encoding="utf-8")
    out_file = tmp_path / "output.txt"
    assert main([str(valid_file), str(out_file)]) == 0
    assert out_file.exists()
    assert "Sample Good CV" in out_file.read_text(encoding="utf-8")
