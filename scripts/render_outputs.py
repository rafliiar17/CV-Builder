from __future__ import annotations

import argparse
import html
import re
import sys
from pathlib import Path

import markdown
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from weasyprint import HTML

CSS = """
@page { size: A4; margin: 16mm; }
body {
  font-family: Arial, sans-serif;
  font-size: 10.5pt;
  line-height: 1.42;
  color: #1f2933;
}
h1 { font-size: 20pt; margin: 0 0 6px; }
h2 { font-size: 13pt; margin: 18px 0 6px; border-bottom: 1px solid #d0d7de; padding-bottom: 3px; }
h3 { font-size: 11.5pt; margin: 14px 0 4px; }
h4 { font-size: 10.5pt; margin: 10px 0 3px; }
p { margin: 0 0 7px; }
ul, ol { margin: 4px 0 8px 18px; padding: 0; }
li { margin-bottom: 3px; }
table { width: 100%; border-collapse: collapse; margin: 8px 0 12px; }
th, td { border: 1px solid #d0d7de; padding: 5px 6px; vertical-align: top; }
th { background: #f6f8fa; }
a { color: #0969da; text-decoration: none; }
code { background: #f6f8fa; padding: 2px 4px; border-radius: 3px; font-family: 'Courier New', monospace; font-size: 9.5pt; }
pre { background: #f6f8fa; padding: 8px 12px; border-radius: 4px; overflow-x: auto; font-family: 'Courier New', monospace; font-size: 9pt; }
blockquote { border-left: 3px solid #d0d7de; padding-left: 10px; color: #57606a; margin: 6px 0; }
hr { border: none; border-top: 1px solid #d0d7de; margin: 16px 0; }
"""

INLINE_MD_RE = re.compile(
    r"(\*\*\*[^*]+?\*\*\*|\*\*[^*]+?\*\*|\*[^*]+?\*|__[^_]+?__|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))"
)


def strip_md(text: str) -> str:
    """Utility to strip markdown syntax to plain text."""
    text = re.sub(r"\*\*\*(.*?)\*\*\*", r"\1", text)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"__(.*?)__", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"`(.*?)`", r"\1", text)
    text = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1", text)
    return text.strip()


def sanitize_xml(text: str) -> str:
    """Remove control characters forbidden in XML 1.0 (keeping tab, newline, carriage return)."""
    return re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", text)


def add_formatted_runs(
    para: Paragraph,
    text: str,
    default_bold: bool = False,
    default_italic: bool = False,
) -> None:
    """Add runs to a python-docx paragraph preserving bold, italic, code, and links."""
    if not text:
        return

    text = sanitize_xml(text)
    parts = INLINE_MD_RE.split(text)
    for part in parts:
        if not part:
            continue

        if part.startswith("***") and part.endswith("***") and len(part) >= 6:
            run = para.add_run(part[3:-3])
            run.bold = True
            run.italic = True
        elif (part.startswith("**") and part.endswith("**") and len(part) >= 4) or (
            part.startswith("__") and part.endswith("__") and len(part) >= 4
        ):
            run = para.add_run(part[2:-2])
            run.bold = True
            if default_italic:
                run.italic = True
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            run = para.add_run(part[1:-1])
            run.italic = True
            if default_bold:
                run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = para.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
            if default_bold:
                run.bold = True
            if default_italic:
                run.italic = True
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            m = re.match(r"^\[(.*?)\]\((.*?)\)$", part)
            if m:
                label = m.group(1)
                run = para.add_run(label)
                run.font.color.rgb = RGBColor(9, 105, 218)
                run.underline = True
                if default_bold:
                    run.bold = True
                if default_italic:
                    run.italic = True
            else:
                run = para.add_run(part)
                if default_bold:
                    run.bold = True
                if default_italic:
                    run.italic = True
        else:
            run = para.add_run(part)
            if default_bold:
                run.bold = True
            if default_italic:
                run.italic = True


def parse_table_row(raw: str) -> list[str]:
    """Parse a markdown table row into cell strings, handling escaped pipes."""
    clean = raw.strip().replace(r"\|", "__PIPE_ESC__")
    if clean.startswith("|"):
        clean = clean[1:]
    if clean.endswith("|"):
        clean = clean[:-1]
    return [c.strip().replace("__PIPE_ESC__", "|") for c in clean.split("|")]


parse_markdown_table_row = parse_table_row


def is_markdown_table_separator(cells_or_raw: list[str] | str) -> bool:
    """Check if a line or parsed cells list represents a markdown table separator row like |---|---|."""
    if isinstance(cells_or_raw, str):
        cells = parse_table_row(cells_or_raw)
    else:
        cells = cells_or_raw
    return bool(cells) and all(bool(re.match(r"^:?-+:?$", c.strip())) for c in cells)


is_table_separator = is_markdown_table_separator


def add_table_to_docx(document: Document, table_lines: list[str]) -> None:
    """Construct a native Word table with Table Grid styling, borders, and light grey header shading."""
    sep_idx = -1
    header_idx = -1
    for idx, line in enumerate(table_lines):
        if is_table_separator(line):
            sep_idx = idx
            header_idx = idx - 1
            break

    if sep_idx != -1 and header_idx >= 0:
        header_cells = parse_table_row(table_lines[header_idx])
        data_lines = [
            table_lines[i]
            for i in range(len(table_lines))
            if i != header_idx and i != sep_idx and not is_table_separator(table_lines[i])
        ]
        data_rows = [parse_table_row(row_line) for row_line in data_lines]
    else:
        header_cells = parse_table_row(table_lines[0])
        data_rows = [parse_table_row(row_line) for row_line in table_lines[1:] if not is_table_separator(row_line)]

    num_cols = max(len(header_cells), max((len(r) for r in data_rows), default=0))
    if num_cols == 0:
        return

    header_cells += [""] * (num_cols - len(header_cells))
    total_rows = 1 + len(data_rows)
    table = document.add_table(rows=total_rows, cols=num_cols)
    table.style = "Table Grid"

    # Header row with light grey shading and bold text
    hdr_row = table.rows[0]
    for col_idx, cell_text in enumerate(header_cells):
        cell = hdr_row.cells[col_idx]
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F2F2"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, cell_text, default_bold=True)

    # Data rows with formatting preserved
    for row_idx, row_data in enumerate(data_rows):
        row = table.rows[row_idx + 1]
        row_data += [""] * (num_cols - len(row_data))
        for col_idx, cell_text in enumerate(row_data):
            cell = row.cells[col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, cell_text)


def render_docx(md_path: Path, out_path: Path | None = None) -> Path:
    """Render markdown file to ATS-compliant DOCX with table and rich inline styling support."""
    if out_path is None:
        out_path = md_path.with_suffix(".docx")
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for h_name in ["Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        if h_name in styles:
            styles[h_name].font.name = "Arial"
            styles[h_name].font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i].rstrip()
        stripped = raw.strip()

        if not stripped:
            i += 1
            continue

        # Fenced code block
        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < n and lines[i].strip().startswith("```"):
                i += 1
            code_text = "\n".join(code_lines)
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(2)
            r = p.add_run(code_text)
            r.font.name = "Courier New"
            r.font.size = Pt(9)
            continue

        # Table block
        if stripped.startswith("|") and "|" in stripped[1:]:
            table_lines = []
            while i < n and lines[i].strip().startswith("|") and "|" in lines[i].strip()[1:]:
                table_lines.append(lines[i].strip())
                i += 1
            add_table_to_docx(document, table_lines)
            continue

        # Headings
        if raw.startswith("# "):
            h = document.add_heading(level=1)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(4)
            add_formatted_runs(h, raw[2:].strip())
        elif raw.startswith("## "):
            h = document.add_heading(level=2)
            h.paragraph_format.space_before = Pt(7)
            h.paragraph_format.space_after = Pt(3)
            add_formatted_runs(h, raw[3:].strip())
        elif raw.startswith("### "):
            h = document.add_heading(level=3)
            h.paragraph_format.space_before = Pt(5)
            h.paragraph_format.space_after = Pt(2)
            add_formatted_runs(h, raw[4:].strip())
        elif raw.startswith("#### "):
            h = document.add_heading(level=4)
            h.paragraph_format.space_before = Pt(4)
            h.paragraph_format.space_after = Pt(2)
            add_formatted_runs(h, raw[5:].strip())
        elif stripped.startswith("- ") or stripped.startswith("* "):
            bullet_text = stripped[2:].strip()
            indent_level = (len(raw) - len(raw.lstrip())) // 2
            p = document.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if indent_level > 0:
                p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            add_formatted_runs(p, bullet_text)
        elif re.match(r"^\d+\.\s+", stripped):
            m = re.match(r"^\d+\.\s+(.*)$", stripped)
            item_text = m.group(1).strip() if m else stripped
            p = document.add_paragraph(style="List Number")
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, item_text)
        elif stripped.startswith(">"):
            quote_text = stripped.lstrip("> ").strip()
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, quote_text, default_italic=True)
        elif stripped in ("---", "***", "___"):
            p = document.add_paragraph("─" * 50)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
        else:
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, stripped)

        i += 1

    document.save(out_path)
    return out_path


def render_pdf(md_path: Path, out_path: Path | None = None) -> Path:
    """Render markdown file to styled A4 PDF via WeasyPrint."""
    if out_path is None:
        out_path = md_path.with_suffix(".pdf")
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    body = markdown.markdown(
        md_path.read_text(encoding="utf-8"),
        extensions=["tables", "sane_lists", "fenced_code"],
        output_format="html5",
    )
    title = html.escape(md_path.stem)
    page = f"<!doctype html><html><head><meta charset='utf-8'><title>{title}</title><style>{CSS}</style></head><body>{body}</body></html>"
    HTML(string=page, base_url=str(md_path.parent)).write_pdf(out_path)
    return out_path


def find_markdown_files(target_dir: Path) -> list[Path]:
    """Find all candidate output markdown files to render recursively, skipping inputs and scratch."""
    if not target_dir.is_dir():
        return []

    md_files: list[Path] = []
    for path in sorted(target_dir.rglob("*.md")):
        parts = set(path.relative_to(target_dir).parts)
        if path.name.startswith("."):
            continue
        if target_dir.name not in ("input", "scratch") and ("input" in parts or "scratch" in parts):
            continue
        md_files.append(path)
    return md_files


def render_file(
    md_path: Path,
    out_docx: Path | None = None,
    out_pdf: Path | None = None,
    fmt: str = "all",
) -> tuple[Path | None, Path | None]:
    """Render a single markdown file into docx and/or pdf."""
    res_docx = None
    res_pdf = None
    if fmt in ("all", "docx"):
        res_docx = render_docx(md_path, out_docx)
    if fmt in ("all", "pdf"):
        res_pdf = render_pdf(md_path, out_pdf)
    return res_docx, res_pdf


def main(argv: list[str] | None = None) -> int:
    """CLI Entry point for rendering markdown outputs."""
    if argv is None:
        argv = sys.argv[1:]

    parser = argparse.ArgumentParser(
        description="Render Markdown outputs to ATS-compliant DOCX and styled PDF."
    )
    parser.add_argument(
        "paths",
        nargs="*",
        help="Markdown file(s) or run directories to render.",
    )
    parser.add_argument(
        "--all",
        dest="all_dir",
        help="Batch render all markdown files under the specified directory.",
    )
    parser.add_argument(
        "--run-dir",
        dest="run_dir",
        help="Candidate run directory to batch render.",
    )
    parser.add_argument(
        "--output-dir",
        dest="output_dir",
        help="Optional destination directory for rendered files.",
    )
    parser.add_argument(
        "--format",
        dest="fmt",
        choices=["all", "docx", "pdf"],
        default="all",
        help="Format to generate: docx, pdf, or all (default: all).",
    )

    args = parser.parse_args(argv)

    targets: list[tuple[Path, Path | None, Path | None]] = []
    out_dir_path = Path(args.output_dir) if args.output_dir else None

    # Handle --all or --run-dir batch options
    batch_dir = args.all_dir or args.run_dir
    if batch_dir:
        b_path = Path(batch_dir)
        if not b_path.exists():
            print(f"error: batch directory not found: {b_path}", file=sys.stderr)
            return 1
        found = find_markdown_files(b_path)
        if not found:
            print(f"warning: no renderable markdown files found under {b_path}", file=sys.stderr)
            return 0
        for f in found:
            if out_dir_path:
                rel = f.relative_to(b_path)
                dest = out_dir_path / rel
                targets.append((f, dest.with_suffix(".docx"), dest.with_suffix(".pdf")))
            else:
                targets.append((f, None, None))

    # Handle positional paths
    for item in args.paths:
        p = Path(item)
        if not p.exists():
            print(f"missing: {p}", file=sys.stderr)
            return 1
        if p.is_dir():
            found = find_markdown_files(p)
            for f in found:
                if out_dir_path:
                    rel = f.relative_to(p)
                    dest = out_dir_path / rel
                    targets.append((f, dest.with_suffix(".docx"), dest.with_suffix(".pdf")))
                else:
                    targets.append((f, None, None))
        elif p.is_file():
            if out_dir_path:
                dest = out_dir_path / p.name
                targets.append((p, dest.with_suffix(".docx"), dest.with_suffix(".pdf")))
            else:
                targets.append((p, None, None))

    if not targets:
        parser.print_usage(file=sys.stderr)
        return 2

    # Deduplicate targets by source md path while keeping order
    seen: set[Path] = set()
    unique_targets: list[tuple[Path, Path | None, Path | None]] = []
    for md_p, out_d, out_p in targets:
        if md_p not in seen:
            seen.add(md_p)
            unique_targets.append((md_p, out_d, out_p))

    failures = 0
    success_count = 0
    for md_path, target_docx, target_pdf in unique_targets:
        try:
            d_path, p_path = render_file(md_path, target_docx, target_pdf, fmt=args.fmt)
            if d_path:
                print(f"rendered: {d_path}")
            if p_path:
                print(f"rendered: {p_path}")
            success_count += 1
        except Exception as e:
            failures += 1
            print(f"error: failed to render {md_path}: {e}", file=sys.stderr)

    if len(unique_targets) > 1:
        print(f"Batch completed: {success_count} succeeded, {failures} failed.")

    return 1 if failures > 0 else 0


if __name__ == "__main__":
    raise SystemExit(main())
