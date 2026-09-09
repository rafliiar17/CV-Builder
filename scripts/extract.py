#!/usr/bin/env python3
"""
CV Brainstormer — CV Extractor Script
Extracts text from PDF, DOCX, DOC, or TXT/MD files into plain text for Agent 00.
Preserves hyperlinks behind text (LinkedIn, GitHub, Portfolios) from PDF annotations and DOCX relationships.

Usage:
    python3 scripts/extract.py <input_file> [output_file]
    python3 scripts/extract.py --help

Examples:
    python3 scripts/extract.py input/my-cv.pdf
    python3 scripts/extract.py input/my-cv.docx input/extracted.txt
    python3 scripts/extract.py output/candidates/rafli-arraafi/2026-06-01/input/original-cv.pdf
"""

import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path


def clean_anchor_text(raw_anchor: str) -> str:
    """Clean extracted anchor text by stripping punctuation and leading/trailing conjunctions."""
    if not raw_anchor:
        return ""
    anchor = raw_anchor.strip(" \t\n\r|:,.;()[]{}<>\"'")
    # Strip accidental conjunctions like 'and GitHub' or 'or Portfolio'
    anchor = re.sub(r"^(and|or|&|\|)\s+", "", anchor, flags=re.IGNORECASE)
    anchor = re.sub(r"\s+(and|or|&|\|)$", "", anchor, flags=re.IGNORECASE)
    return anchor.strip(" \t\n\r|:,.;()[]{}<>\"'")


def extract_pdf(filepath: str) -> str:
    """Extract text from PDF using pdfplumber, preserving hyperlinks."""
    try:
        import pdfplumber
    except ImportError:
        print("ERROR: pdfplumber not installed.", file=sys.stderr)
        print("Run: pip install pdfplumber --break-system-packages", file=sys.stderr)
        sys.exit(1)

    text_parts = []
    total_links_found = 0

    with pdfplumber.open(filepath) as pdf:
        num_pages = len(pdf.pages)
        print(f"PDF has {num_pages} page(s).")
        for i, page in enumerate(pdf.pages, 1):
            page_text = page.extract_text() or ""

            # Extract hyperlinks if available
            hyperlinks = getattr(page, "hyperlinks", None) or []
            words = page.extract_words() if hasattr(page, "extract_words") else []

            page_links = []
            seen_uris = set()

            for h in hyperlinks:
                uri = h.get("uri")
                if not uri or not isinstance(uri, str):
                    continue
                uri = uri.strip()
                if not uri or uri in seen_uris:
                    continue
                seen_uris.add(uri)

                # Bounding box of the hyperlink
                x0 = min(h.get("x0", 0), h.get("x1", 0))
                x1 = max(h.get("x0", 0), h.get("x1", 0))
                top = min(h.get("top", 0), h.get("bottom", 0))
                bottom = max(h.get("top", 0), h.get("bottom", 0))

                # Find matching words that overlap significantly with the annotation bbox
                tol = 2.0
                overlapping_words = []
                for w in words:
                    wx0 = w.get("x0", 0)
                    wx1 = w.get("x1", 0)
                    wtop = w.get("top", 0)
                    wbot = w.get("bottom", 0)

                    # Compute overlap area
                    inter_x0 = max(wx0, x0 - tol)
                    inter_x1 = min(wx1, x1 + tol)
                    inter_top = max(wtop, top - tol)
                    inter_bot = min(wbot, bottom + tol)

                    if inter_x1 > inter_x0 and inter_bot > inter_top:
                        inter_area = (inter_x1 - inter_x0) * (inter_bot - inter_top)
                        word_area = max(0.1, (wx1 - wx0) * (wbot - wtop))
                        # Match if at least 35% of word is inside link bbox
                        if (inter_area / word_area) >= 0.35:
                            overlapping_words.append(w["text"])

                clean_anchor = clean_anchor_text(" ".join(overlapping_words))

                page_links.append((clean_anchor, uri))
                total_links_found += 1

                # Embed markdown link into page text if clean_anchor exists in page_text
                # and uri is not already in page_text
                if clean_anchor and len(clean_anchor) >= 2 and uri not in page_text:
                    if clean_anchor in page_text:
                        if f"[{clean_anchor}]" not in page_text and clean_anchor != uri:
                            pattern = re.compile(r"\b" + re.escape(clean_anchor) + r"\b")
                            if pattern.search(page_text):
                                page_text = pattern.sub(f"[{clean_anchor}]({uri})", page_text, count=1)
                            else:
                                page_text = page_text.replace(clean_anchor, f"[{clean_anchor}]({uri})", 1)

            if page_text.strip():
                text_parts.append(f"--- PAGE {i} ---")
                text_parts.append(page_text.strip())

                # Append detected links on this page
                if page_links:
                    link_lines = []
                    for anchor, uri in page_links:
                        if anchor and anchor != uri:
                            link_lines.append(f"- [{anchor}]({uri})")
                        else:
                            link_lines.append(f"- <{uri}>")
                    text_parts.append("\n[Detected Links on Page " + str(i) + "]\n" + "\n".join(link_lines))
            else:
                print(f"WARNING: Page {i} has no extractable text (may be image-based).", file=sys.stderr)
                print("Consider using OCR for scanned PDFs.", file=sys.stderr)
                if page_links:
                    text_parts.append(f"--- PAGE {i} (IMAGE / NO TEXT) ---")
                    link_lines = [f"- <{uri}>" for _, uri in page_links]
                    text_parts.append("[Detected Links on Page " + str(i) + "]\n" + "\n".join(link_lines))

    if total_links_found > 0:
        print(f"Extracted {total_links_found} hyperlink(s) from PDF annotations.")

    return "\n\n".join(text_parts)


def extract_docx_paragraph(p, part) -> str:
    """Extract text from a DOCX paragraph preserving hyperlinks from relationships."""
    from docx.oxml.ns import qn

    chunks = []
    for elem in p._p:
        tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
        if tag == "r":
            text = "".join(node.text for node in elem.iter() if node.tag.endswith("t") and node.text)
            chunks.append(text)
        elif tag == "hyperlink":
            r_id = elem.attrib.get(qn("r:id")) or elem.attrib.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            )
            url = None
            if r_id and part and r_id in part.rels:
                rel = part.rels[r_id]
                url = getattr(rel, "target_ref", None)

            link_text = "".join(node.text for node in elem.iter() if node.tag.endswith("t") and node.text)
            link_text_clean = link_text.strip()

            if url and link_text_clean:
                if url.strip() == link_text_clean or link_text_clean.startswith("http://") or link_text_clean.startswith("https://"):
                    chunks.append(link_text)
                else:
                    chunks.append(f"[{link_text}]({url})")
            elif url:
                chunks.append(f"<{url}>")
            elif link_text:
                chunks.append(link_text)
        else:
            # Fallback for nested or other inline elements
            text = "".join(node.text for node in elem.iter() if node.tag.endswith("t") and node.text)
            if text:
                chunks.append(text)
    return "".join(chunks)


def extract_docx(filepath: str) -> str:
    """Extract text from DOCX using python-docx, preserving hyperlinks, headers, and tables."""
    try:
        from docx import Document
    except ImportError:
        print("ERROR: python-docx not installed.", file=sys.stderr)
        print("Run: pip install python-docx --break-system-packages", file=sys.stderr)
        sys.exit(1)

    doc = Document(filepath)
    text_parts = []
    part = getattr(doc, "part", None)

    # Check header paragraphs (resumes often put contact info and links in headers)
    for section in doc.sections:
        header = getattr(section, "header", None)
        if header and hasattr(header, "paragraphs"):
            h_part = getattr(header, "part", part)
            for p in header.paragraphs:
                p_text = extract_docx_paragraph(p, h_part).strip()
                if p_text and p_text not in text_parts:
                    text_parts.append(p_text)

    # Extract body paragraphs
    for para in doc.paragraphs:
        p_text = extract_docx_paragraph(para, part).strip()
        if p_text:
            text_parts.append(p_text)

    # Extract text from tables with hyperlinks preserved
    for table in doc.tables:
        for row in table.rows:
            cell_texts = []
            for cell in row.cells:
                paras = [extract_docx_paragraph(p, part).strip() for p in cell.paragraphs]
                cell_content = " ".join(p for p in paras if p).strip()
                if cell_content:
                    cell_texts.append(cell_content)
            if cell_texts:
                text_parts.append(" | ".join(cell_texts))

    # Check if any external hyperlinks from document relationships were not captured
    all_text = "\n".join(text_parts)
    unseen_links = []
    if part and hasattr(part, "rels"):
        for rel_id, rel in part.rels.items():
            if getattr(rel, "is_external", False):
                rel_type = getattr(rel, "reltype", "")
                if "hyperlink" in rel_type.lower():
                    target = getattr(rel, "target_ref", None)
                    if target and target not in all_text:
                        unseen_links.append(target)

    if unseen_links:
        text_parts.append("\n[Detected Document Links]")
        for link_url in unseen_links:
            text_parts.append(f"- <{link_url}>")

    return "\n".join(text_parts)


def extract_doc(filepath: str) -> str:
    """Handle legacy .doc files with conversion or clear helpful error message."""
    # 1. Check if it's actually an OOXML file (docx misnamed as doc)
    if zipfile.is_zipfile(filepath):
        try:
            return extract_docx(filepath)
        except Exception:
            pass

    # 2. Check if headless LibreOffice or soffice is available for conversion
    converter = shutil.which("libreoffice") or shutil.which("soffice")
    if converter:
        with tempfile.TemporaryDirectory() as temp_dir:
            try:
                res = subprocess.run(
                    [converter, "--headless", "--convert-to", "docx", str(filepath), "--outdir", temp_dir],
                    capture_output=True,
                    text=True,
                    timeout=45,
                )
                if res.returncode == 0:
                    converted_files = list(Path(temp_dir).glob("*.docx"))
                    if converted_files:
                        print(f"NOTE: Converted legacy .doc to .docx via {Path(converter).name} for extraction.")
                        return extract_docx(str(converted_files[0]))
            except Exception as e:
                print(f"WARNING: Automatic conversion via {converter} failed: {e}", file=sys.stderr)

    # 3. If conversion failed or converter not found, provide a clear, helpful error message
    print(f"\nERROR: Cannot read legacy binary .doc format directly: {filepath}", file=sys.stderr)
    print("python-docx only supports the modern .docx format (Office Open XML).", file=sys.stderr)
    print("\nPlease convert the file to .docx or .pdf before extracting:", file=sys.stderr)
    print(f"  Option 1 (LibreOffice CLI): libreoffice --headless --convert-to docx \"{filepath}\"", file=sys.stderr)
    print("  Option 2: Open in Microsoft Word, LibreOffice Writer, or Google Docs and 'Save As...' .docx or .pdf", file=sys.stderr)
    sys.exit(1)


def extract_txt(filepath: str) -> str:
    """Read plain text or markdown file."""
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def extract_file(filepath: str | Path, output_path: str | Path | None = None) -> str:
    """Extract text from a supported file format (.pdf, .docx, .doc, .txt, .md)."""
    in_path = Path(filepath)
    if not in_path.exists():
        raise FileNotFoundError(f"File not found: {in_path}")

    ext = in_path.suffix.lower()
    if ext == ".pdf":
        text = extract_pdf(str(in_path))
    elif ext == ".docx":
        text = extract_docx(str(in_path))
    elif ext == ".doc":
        text = extract_doc(str(in_path))
    elif ext in (".txt", ".md"):
        text = extract_txt(str(in_path))
    else:
        raise ValueError(f"Unsupported file type: '{ext}'. Supported formats: .pdf, .docx, .doc, .txt, .md")

    if output_path is not None:
        out_p = Path(output_path)
        out_p.parent.mkdir(parents=True, exist_ok=True)
        with open(out_p, "w", encoding="utf-8") as f:
            f.write(f"=== EXTRACTED FROM: {in_path.name} ===\n\n")
            f.write(text)
            f.write("\n\n=== END OF EXTRACTION ===\n")

    return text


def main(argv: list[str] | None = None) -> int:
    args = argv if argv is not None else sys.argv[1:]
    if len(args) < 1 or args[0] in ("-h", "--help"):
        print(__doc__.strip())
        return 0 if len(args) >= 1 and args[0] in ("-h", "--help") else 1

    input_path = Path(args[0])
    output_path = Path(args[1]) if len(args) >= 2 else Path("input/extracted.txt")

    if not input_path.exists():
        print(f"ERROR: File not found: {input_path}", file=sys.stderr)
        return 1

    try:
        text = extract_file(input_path, output_path)
        print(f"Extracting: {input_path} ({input_path.suffix.lower()})")
        print(f"Done. Output saved to: {output_path}")
        print(f"Character count: {len(text):,}")
        print(f"Word count (approx): {len(text.split()):,}")
        print()
        print("Next step: Pass this file to Agent 00 (Extractor) for normalization.")
        return 0
    except (FileNotFoundError, ValueError) as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
